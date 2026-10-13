
from pptx import Presentation

def load_ppt(file_path):
    prs = Presentation(file_path)
    slide_text = []

    for slide_number, slide in enumerate(prs.slides, start=1):
        content = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                content.append(shape.text)

            if shape.has_table:
                for row in shape.table.rows:
                    content.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

        if content:
            slide_text.append(
                f"Slide {slide_number}\n" + "\n".join(content)
            )

    return "\n\n".join(slide_text)
